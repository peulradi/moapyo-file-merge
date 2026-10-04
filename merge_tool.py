"""Offline spreadsheet consolidation with row provenance and explicit review.

This is product runtime code. It does not call an AI API, network, or Excel.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import uuid
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

VERSION = "0.1.0"
META = ["__원본파일", "__원본시트", "__원본행", "__검수"]
MAX_FILES = 30
MAX_ROWS = 100_000
MAX_COLS = 100
MAX_FILE_BYTES = 25 * 1024 * 1024


class InputError(ValueError):
    """An input problem that must be fixed before any output is delivered."""


@dataclass
class Record:
    values: dict[str, Any]
    source: str
    sheet: str
    row: int
    issues: set[str] = field(default_factory=set)


@dataclass
class Result:
    columns: list[str]
    records: list[Record]
    excluded: list[Record]
    report: dict[str, Any]


def load_config(path: str | Path | None) -> dict[str, Any]:
    config = {
        "keys": [], "aliases": {}, "required": [], "numeric_columns": [],
        "dedupe_exact": False, "header_row": 1, "sheet": None,
    }
    if path:
        supplied = json.loads(Path(path).read_text(encoding="utf-8-sig"))
        if not isinstance(supplied, dict):
            raise InputError("설정은 JSON 객체여야 합니다.")
        unknown = set(supplied) - set(config)
        if unknown:
            raise InputError(f"알 수 없는 설정: {', '.join(sorted(unknown))}")
        config.update(supplied)
    for key in ("keys", "required", "numeric_columns"):
        if not isinstance(config[key], list) or not all(isinstance(x, str) and x for x in config[key]):
            raise InputError(f"{key}는 비어 있지 않은 열 이름의 배열이어야 합니다.")
        if len(config[key]) != len(set(config[key])):
            raise InputError(f"{key}에 중복된 열 이름이 있습니다.")
    if not isinstance(config["aliases"], dict) or not all(
        isinstance(k, str) and k and isinstance(v, str) and v
        for k, v in config["aliases"].items()
    ):
        raise InputError("aliases는 원래 열 이름과 새 열 이름의 대응표여야 합니다.")
    if type(config["dedupe_exact"]) is not bool:
        raise InputError("dedupe_exact는 true 또는 false여야 합니다.")
    if type(config["header_row"]) is not int or not 1 <= config["header_row"] <= 100:
        raise InputError("header_row는 1~100 정수여야 합니다.")
    if config["sheet"] is not None and not isinstance(config["sheet"], str):
        raise InputError("sheet는 시트 이름 또는 null이어야 합니다.")
    if set(config["keys"]) & set(config["numeric_columns"]):
        raise InputError("식별키는 숫자로 변환할 수 없습니다. 선행 0을 보존하세요.")
    return config


def normalize_headers(raw: list[Any], aliases: dict[str, str], source: str) -> list[str]:
    if not raw or len(raw) > MAX_COLS:
        raise InputError(f"{source}: 열 개수가 1~{MAX_COLS} 범위를 벗어났습니다.")
    headers = []
    for cell in raw:
        name = "" if cell is None else str(cell).strip()
        if not name:
            raise InputError(f"{source}: 제목 행에 빈 열 이름이 있습니다.")
        name = aliases.get(name, name)
        if name in META or name.startswith("__"):
            raise InputError(f"{source}: {name}은 추적 기록용 예약 열 이름입니다.")
        headers.append(name)
    if len(set(headers)) != len(headers):
        raise InputError(f"{source}: 열 이름 또는 별칭이 중복됩니다.")
    return headers


def read_csv(path: Path, config: dict[str, Any]) -> tuple[list[str], list[Record], str, int]:
    raw = path.read_bytes()
    text = None
    encoding = None
    for candidate in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise InputError(f"{path.name}: UTF-8 또는 CP949로 읽을 수 없습니다.")
    if "\x00" in text:
        raise InputError(f"{path.name}: NUL 문자가 있습니다. UTF-8 CSV로 다시 저장하세요.")
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
    headers = None
    records = []
    blank_rows = 0
    prior_line = 0
    for row_number, row in enumerate(reader, start=1):
        source_line = prior_line + 1
        prior_line = reader.line_num
        if row_number < config["header_row"]:
            continue
        if headers is None:
            headers = normalize_headers(row, config["aliases"], path.name)
            continue
        if not row or all(cell == "" for cell in row):
            blank_rows += 1
            continue
        if len(row) > len(headers):
            raise InputError(f"{path.name}: {row_number}행의 값 개수가 열 이름보다 많습니다.")
        padded = row + [""] * (len(headers) - len(row))
        records.append(Record(dict(zip(headers, padded)), path.name, "CSV", source_line))
        if len(records) > MAX_ROWS:
            raise InputError("처리 행 수 제한을 넘었습니다.")
    if headers is None:
        raise InputError(f"{path.name}: 제목 행이 없습니다.")
    return headers, records, str(encoding), blank_rows


def read_xlsx(path: Path, config: dict[str, Any]) -> tuple[list[str], list[Record], str, int]:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
            raise InputError(f"{path.name}: 압축 해제 데이터가 안전 처리 한도를 넘습니다.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for candidate_sheet in workbook.worksheets:
            if candidate_sheet.max_row is None or candidate_sheet.max_column is None:
                candidate_sheet.calculate_dimension(force=True)
        if config["sheet"]:
            if config["sheet"] not in workbook.sheetnames:
                raise InputError(f"{path.name}: {config['sheet']} 시트가 없습니다.")
            sheet = workbook[config["sheet"]]
        else:
            nonempty = [s for s in workbook.worksheets if s.max_row > 1]
            if len(nonempty) > 1:
                raise InputError(f"{path.name}: 데이터 시트가 여러 개입니다. 설정에 sheet를 지정하세요.")
            sheet = nonempty[0] if nonempty else workbook.worksheets[0]
        if (sheet.max_column or 0) > MAX_COLS or (sheet.max_row or 0) > MAX_ROWS + 100:
            raise InputError(f"{path.name}: 데이터 크기 제한을 넘었습니다.")
        records = []
        headers = None
        blank_rows = 0
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if row_number < config["header_row"]:
                continue
            values = [cell.value for cell in cells]
            if headers is None:
                while values and values[-1] is None:
                    values.pop()
                headers = normalize_headers(values, config["aliases"], path.name)
                continue
            if all(v is None or v == "" for v in values):
                blank_rows += 1
                continue
            if any(v is not None for v in values[len(headers):]):
                raise InputError(f"{path.name}: {row_number}행에 제목 없는 데이터가 있습니다.")
            row = Record(dict(zip(headers, values[:len(headers)])), path.name, sheet.title, row_number)
            for column, cell in zip(headers, cells):
                if cell.data_type == "f":
                    row.issues.add(f"수식원문:{column}")
                if column in config["keys"] and isinstance(cell.value, (int, float)):
                    row.issues.add(f"숫자형식별자확인:{column}")
            records.append(row)
        if headers is None:
            raise InputError(f"{path.name}: 제목 행이 없습니다.")
        return headers, records, "XLSX", blank_rows
    finally:
        workbook.close()


def signature(values: list[Any]) -> str:
    # Types matter: the ID '001' must never collapse into numeric 1.
    parts = [(type(v).__name__, v.isoformat() if isinstance(v, (datetime, date)) else v) for v in values]
    return json.dumps(parts, ensure_ascii=False, default=str, sort_keys=True)


def empty(value: Any) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def numeric(value: Any) -> int | float:
    if isinstance(value, bool):
        raise InputError("참/거짓 값은 숫자가 아닙니다.")
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise InputError("유한한 숫자가 아닙니다.")
        return value
    text = str(value).strip()
    if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text):
        raise InputError("숫자 형식을 확인하세요.")
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise InputError("숫자 형식을 확인하세요.") from exc
    if abs(number) > Decimal("1e14"):
        raise InputError("큰 숫자는 정밀도 손실 방지를 위해 원문으로 보존합니다.")
    return int(number) if number == number.to_integral_value() else float(number)


def consolidate(paths: list[Path], config: dict[str, Any]) -> Result:
    if set(config["keys"]) & set(config["numeric_columns"]):
        raise InputError("식별키는 숫자로 변환할 수 없습니다. 선행 0을 보존하세요.")
    if not paths or len(paths) > MAX_FILES:
        raise InputError(f"파일을 1~{MAX_FILES}개 선택하세요.")
    resolved = [p.resolve(strict=True) for p in paths]
    if len(set(resolved)) != len(resolved):
        raise InputError("같은 파일을 두 번 선택했습니다.")
    if len({p.name.casefold() for p in resolved}) != len(resolved):
        raise InputError("서로 다른 폴더의 파일명이 같습니다. 추적 가능하도록 이름을 변경하세요.")
    columns: list[str] = []
    records = []
    files = []
    blank_rows = 0
    for path in resolved:
        if path.stat().st_size > MAX_FILE_BYTES:
            raise InputError(f"{path.name}: 25MB 이하 파일만 처리합니다.")
        suffix = path.suffix.lower()
        if suffix not in (".csv", ".tsv", ".xlsx"):
            raise InputError(f"{path.name}: CSV, TSV, XLSX 파일만 지원합니다.")
        if suffix == ".xlsx":
            headers, new, encoding, blanks = read_xlsx(path, config)
        else:
            headers, new, encoding, blanks = read_csv(path, config)
        required = set(config["keys"] + config["required"] + config["numeric_columns"])
        missing = required - set(headers)
        if missing:
            raise InputError(f"{path.name}: 필요한 열이 없습니다: {', '.join(sorted(missing))}")
        for header in headers:
            if header not in columns:
                columns.append(header)
        records.extend(new)
        if len(records) > MAX_ROWS or len(columns) > MAX_COLS:
            raise InputError("전체 데이터 크기 제한을 넘었습니다.")
        blank_rows += blanks
        files.append({"name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                      "rows": len(new), "blank_rows": blanks, "encoding": encoding, "columns": headers})
    exact_seen: dict[str, Record] = {}
    kept = []
    excluded = []
    duplicates = 0
    by_key: dict[str, list[Record]] = defaultdict(list)
    for record in records:
        for col in columns:
            record.values.setdefault(col, None)
        for col in config["required"]:
            if empty(record.values[col]):
                record.issues.add(f"필수값누락:{col}")
        for col in config["keys"]:
            if empty(record.values[col]):
                record.issues.add(f"식별키누락:{col}")
        for col in config["numeric_columns"]:
            if empty(record.values[col]):
                record.issues.add(f"숫자값누락:{col}")
                continue
            try:
                record.values[col] = numeric(record.values[col])
            except InputError:
                record.issues.add(f"숫자형식확인:{col}")
        for col, value in record.values.items():
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
                record.issues.add(f"문자열실행방지:{col}")
        exact = signature([record.values[c] for c in columns])
        if exact in exact_seen:
            duplicates += 1
            record.issues.add("전체값일치중복")
            exact_seen[exact].issues.add("전체값일치중복")
            if config["dedupe_exact"]:
                record.issues.add("선택설정으로제외")
                excluded.append(record)
                continue
        else:
            exact_seen[exact] = record
        kept.append(record)
        if config["keys"] and all(not empty(record.values[c]) for c in config["keys"]):
            by_key[signature([record.values[c] for c in config["keys"]])].append(record)
    conflicts = 0
    for group in by_key.values():
        variants = {signature([record.values[c] for c in columns]) for record in group}
        if len(variants) > 1:
            conflicts += 1
            for record in group:
                record.issues.add("동일키값다름")
    counts = Counter(issue for record in kept for issue in record.issues)
    report = {
        "schema_version": 1, "tool_version": VERSION,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "synthetic_demo": False, "input_files": files, "columns": columns,
        "input_rows": len(records), "output_rows": len(kept), "excluded_rows": len(excluded),
        "blank_rows_skipped": blank_rows, "exact_duplicate_rows": duplicates,
        "conflicting_key_groups": conflicts, "review_rows": sum(bool(r.issues) for r in kept),
        "issue_counts": dict(sorted(counts.items())), "config": config,
        "reconciles": len(records) == len(kept) + len(excluded),
        "limits": {"files": MAX_FILES, "rows": MAX_ROWS, "columns": MAX_COLS, "bytes_per_file": MAX_FILE_BYTES},
        "behavior": "수식은 계산하지 않고 원문 문자열로 저장. 원본 값의 공백/날짜/식별자는 임의 정규화하지 않음.",
    }
    return Result(columns, kept, excluded, report)


def append_literal(sheet, values: list[Any]) -> None:
    for value in values:
        if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", value)):
            raise InputError("엑셀 셀에 저장할 수 없는 길이 또는 제어 문자가 있습니다. 원본을 확인하세요.")
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"
            if value.startswith(("=", "+", "-", "@")):
                cell.quotePrefix = True


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="193B5C")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    for column in sheet.columns:
        letter = column[0].column_letter
        longest = max((len(str(c.value or "")) for c in column[:100]), default=10)
        sheet.column_dimensions[letter].width = min(max(14, longest * 1.4 + 2), 54)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Malgun Gothic", size=10)
            cell.alignment = Alignment(vertical="center")


def export_result(result: Result, destination: Path) -> Path:
    """Reserve a new folder; never replace input files or an existing result."""
    destination = destination.resolve()
    destination.mkdir(parents=False, exist_ok=False)
    workbook = Workbook()
    output = workbook.active
    output.title = "취합결과"
    review = workbook.create_sheet("검수대상")
    removed = workbook.create_sheet("제외기록")
    summary = workbook.create_sheet("처리기록")
    header = result.columns + META
    for sheet in (output, review, removed):
        append_literal(sheet, header)
    for record in result.records:
        row = [record.values[c] for c in result.columns] + [record.source, record.sheet, record.row, "; ".join(sorted(record.issues))]
        append_literal(output, row)
        if record.issues:
            append_literal(review, row)
    for record in result.excluded:
        row = [record.values[c] for c in result.columns] + [record.source, record.sheet, record.row, "; ".join(sorted(record.issues))]
        append_literal(removed, row)
    append_literal(summary, ["항목", "값"])
    for key in ("input_rows", "output_rows", "excluded_rows", "blank_rows_skipped", "exact_duplicate_rows", "conflicting_key_groups", "review_rows", "reconciles"):
        append_literal(summary, [key, result.report[key]])
    append_literal(summary, ["작업 기준", "전체 값 일치 중복만 선택적으로 제외. 식별키 충돌은 유지 후 검수."])
    append_literal(summary, ["수식", "계산하지 않음. 원문을 실행되지 않는 문자열로 보존."])
    for source in result.report["input_files"]:
        append_literal(summary, [source["name"], f"입력 {source['rows']}행 / SHA-256 {source['sha256']}"])
    for sheet in workbook:
        style_sheet(sheet)
    temporary = destination / f".{uuid.uuid4().hex}.xlsx"
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, destination / "취합결과.xlsx")
    (destination / "처리기록.json").write_text(json.dumps(result.report, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination / "취합결과.xlsx"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="엑셀·CSV 취합 및 중복 검수. 네트워크 없이 실행합니다.")
    parser.add_argument("inputs", nargs="*", type=Path)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--version", action="version", version=VERSION)
    args = parser.parse_args(argv)
    if not args.inputs:
        from gui import run_gui
        run_gui()
        return 0
    if not args.out:
        parser.error("명령행 처리에는 --out 새_결과폴더가 필요합니다.")
    try:
        result = consolidate(args.inputs, load_config(args.config))
        output = export_result(result, args.out)
        print(json.dumps({"output": str(output), "input_rows": result.report["input_rows"],
                          "output_rows": result.report["output_rows"], "review_rows": result.report["review_rows"]}, ensure_ascii=False))
        return 0
    except (ValueError, OSError, csv.Error, KeyError) as exc:
        print(f"처리 중단: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
